import csv
import json
import hashlib
import subprocess
import sys
import tempfile
import unittest
import re
from copy import deepcopy
from pathlib import Path

import openpyxl
from finance_close.engine import FIELDS, process
from finance_close.io import load_rows, load_catalog
from test_engine import row, control

ROOT = Path(__file__).resolve().parents[1]
CUSTOM = {'centros': {'A': 'Área A'}, 'categorias': {
    'FAT': {'nome': 'Faturamento', 'natureza': 'receita', 'sinal': 'nao_negativo'},
    'CUS': {'nome': 'Custos', 'natureza': 'gasto', 'sinal': 'nao_positivo'}},
    'obrigatorios': ['registro_id', 'data_competencia', 'categoria', 'valor_brl']}
PROFILE = {'nome': 'Meu fechamento', 'entrada': {
    'separador': ';', 'aba': 'Dados', 'colunas': dict(zip(FIELDS, ['ID', 'Data', 'Centro', 'Conta', 'Descrição', 'Valor'])),
    'formato_data': 'br', 'decimal': ',', 'milhar': '.', 'permitir_colunas_extras': False}}


def put_profile_csv(base, values):
    directory = base / 'input'
    directory.mkdir()
    with (directory / 'a.csv').open('w', encoding='utf-8', newline='') as stream:
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(['ID', 'Data', 'Centro', 'Conta', 'Descrição', 'Valor'])
        writer.writerows(values)
    return directory


class ReuseRulesTests(unittest.TestCase):
    def test_renamed_revenue_and_optional_blanks_use_catalog_rules(self):
        result = process([row(id='1', amount='100', centro_custo='A', categoria='FAT', descricao=''),
                          row(id='2', amount='-30', centro_custo='', categoria='CUS', linha=3)], control(2, '70'), '2025-09', CUSTOM)
        self.assertTrue(result['liberado'])
        self.assertEqual(result['resumo']['receitas_centavos'], 10000)
        self.assertEqual(result['resumo']['gastos_centavos'], 3000)
        self.assertEqual(result['resumo']['total_aprovado_centavos'], 7000)

    def test_invalid_category_rules_fail_before_processing(self):
        for nature, sign in [('receitx', 'nao_negativo'), ('receita', 'qualquer')]:
            with self.subTest(nature=nature, sign=sign):
                catalog = deepcopy(CUSTOM)
                catalog['categorias']['FAT'].update(natureza=nature, sinal=sign)
                with self.assertRaises(ValueError):
                    process([row()], control(), '2025-09', catalog)

    def test_structural_fields_cannot_be_disabled(self):
        catalog = deepcopy(CUSTOM)
        catalog['obrigatorios'] = ['descricao']
        with self.assertRaises(ValueError):
            process([], {}, '2025-09', catalog)

    def test_free_sign_allows_explicit_reversal(self):
        catalog = deepcopy(CUSTOM)
        catalog['categorias']['CUS']['sinal'] = 'livre'
        result = process([row(amount='30', centro_custo='A', categoria='CUS')], control(1, '30'), '2025-09', catalog)
        self.assertTrue(result['liberado'])
        self.assertEqual(result['resumo']['gastos_centavos'], -3000)


class ReuseInputTests(unittest.TestCase):
    def test_brazilian_mapping_normalizes_and_preserves_source(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = put_profile_csv(Path(temp), [['0001', '15/09/2025', 'A', 'FAT', '', '1.250,50'], ['0002', '15/09/2025', 'A', 'CUS', 'custo', '-250,50']])
            original = (directory / 'a.csv').read_bytes()
            rows, sources = load_rows(directory, PROFILE)
            self.assertEqual(rows[0]['registro_id'], '0001')
            self.assertEqual(rows[0]['data_competencia'], '2025-09-15')
            self.assertEqual(rows[0]['valor_brl'], '1250.50')
            self.assertEqual(rows[0]['campos_origem']['valor_brl'], '1.250,50')
            self.assertEqual(rows[1]['linha'], 3)
            self.assertEqual(sources['a.csv']['sha256'], hashlib.sha256(original).hexdigest())
            result = process(rows, control(2, '1000'), '2025-09', CUSTOM, sources)
            self.assertTrue(result['liberado'])
            self.assertEqual(result['resumo']['total_aprovado_centavos'], 100000)
            self.assertEqual((directory / 'a.csv').read_bytes(), original)

    def test_wrong_explicit_formats_cannot_be_approved_as_canonical(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = put_profile_csv(Path(temp), [['1', '2025-09-15', 'A', 'FAT', 'texto', '1250.50']])
            rows, sources = load_rows(directory, PROFILE)
            result = process(rows, control(1, '1250.50'), '2025-09', CUSTOM, sources)
            self.assertFalse(result['liberado'])
            self.assertEqual(result['resumo']['valores_nao_quantificaveis'], 1)
            self.assertIn('DATA_FORMATO_INVALIDO', result['linhas'][0]['erros'])

    def test_ambiguous_mapping_and_undeclared_extra_columns_fail(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = put_profile_csv(Path(temp), [['1', '15/09/2025', 'A', 'FAT', 'texto', '1,00']])
            profile = deepcopy(PROFILE)
            profile['entrada']['colunas']['valor_brl'] = 'ID'
            with self.assertRaisesRegex(ValueError, 'coluna|mapeamento'):
                load_rows(directory, profile)
            path = directory / 'a.csv'
            text = path.read_text(encoding='utf-8').replace('Valor\n', 'Valor;Extra\n').replace('1,00\n', '1,00;ignorar\n')
            path.write_text(text, encoding='utf-8')
            with self.assertRaises(ValueError):
                load_rows(directory, PROFILE)
            permissive = deepcopy(PROFILE)
            permissive['entrada']['permitir_colunas_extras'] = True
            rows, _ = load_rows(directory, permissive)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]['valor_brl'], '1.00')

    def test_excel_mapping_keeps_typed_amounts_and_source_sheet(self):
        with tempfile.TemporaryDirectory() as temp:
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = 'Dados'
            sheet.append(['ID', 'Data', 'Centro', 'Conta', 'Descrição', 'Valor'])
            sheet.append(['01', '15/09/2025', 'A', 'FAT', 'Receita', 1250.50])
            book.save(Path(temp) / 'a.xlsx')
            book.close()
            rows, _ = load_rows(temp, PROFILE)
            self.assertEqual(rows[0]['aba'], 'Dados')
            self.assertEqual(rows[0]['valor_brl'], '1250.5')


class ReuseReportTests(unittest.TestCase):
    def test_excel_uses_resolved_nature_for_case_and_numeric_category_codes(self):
        from finance_close.export import write_outputs
        for revenue, expense in [('FAT', 'fat'), ('01', '1')]:
            with self.subTest(revenue=revenue, expense=expense), tempfile.TemporaryDirectory() as temp:
                catalog = deepcopy(CUSTOM)
                catalog['categorias'] = {revenue: CUSTOM['categorias']['FAT'], expense: CUSTOM['categorias']['CUS']}
                result = process([row(id='1', amount='100', categoria=revenue, centro_custo='A'), row(id='2', amount='-30', categoria=expense, centro_custo='A', linha=3)], control(2, '70'), '2025-09', catalog)
                self.assertTrue(result['liberado'])
                write_outputs(result, Path(temp) / 'out')
                book = openpyxl.load_workbook(Path(temp) / 'out/fechamento.xlsx', data_only=False)
                self.assertEqual(book['Base aprovada']['L7'].value, 'receita')
                self.assertEqual(book['Base aprovada']['L8'].value, 'gasto')
                self.assertIn('L7:L8', book['Resumo']['B17'].value)
                self.assertIn('"receita"', book['Resumo']['B17'].value)
                self.assertIn('"gasto"', book['Resumo']['B18'].value)
                book.close()
                cached = openpyxl.load_workbook(Path(temp) / 'out/fechamento.xlsx', data_only=True)
                self.assertEqual([cached['Resumo'][c].value for c in ('B17', 'B18', 'B19', 'B22')], [100, 30, 70, 70])
                cached.close()

    def test_data_tokens_are_not_interpreted_as_template_source(self):
        from finance_close.report import render_run_report
        with tempfile.TemporaryDirectory() as temp:
            result = process([row(categoria='FAT', centro_custo='A', descricao='__JS__ __CSS__ __DATA__')], control(), '2025-09', CUSTOM)
            path = Path(temp) / 'relatorio.html'
            render_run_report(result, path, title='__CSS__')
            html = path.read_text(encoding='utf-8')
            data = json.loads(re.search(r'id="demo-data">(.*?)</script>', html, re.S).group(1))
            self.assertEqual(data['execucao']['linhas'][0]['descricao'], '__JS__ __CSS__ __DATA__')
            self.assertIn('<title>__CSS__</title>', html)

    def test_run_profile_generates_real_html_and_excel_totals(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            directory = put_profile_csv(base, [['0001', '15/02/2026', 'A', 'FAT', '', '1.250,50'], ['0002', '15/02/2026', 'A', 'CUS', 'custo', '-250,50']])
            profile = deepcopy(PROFILE)
            profile['nome'] = 'Equipe <img src=x onerror=alert(1)>'
            (base / 'perfil.json').write_text(json.dumps(profile), encoding='utf-8')
            (base / 'catalogo.json').write_text(json.dumps(CUSTOM), encoding='utf-8')
            (base / 'controle.csv').write_text('arquivo;linhas_esperadas;total_esperado_brl\na.csv;2;1000.00\n', encoding='utf-8')
            command = [sys.executable, '-m', 'finance_close', 'run', '--input', str(directory), '--controls', str(base / 'controle.csv'), '--catalog', str(base / 'catalogo.json'), '--profile', str(base / 'perfil.json'), '--period', '2026-02', '--output']
            for period_args, folder, expected_code in [([], 'out', 0), (['--period', '2026-03'], 'pending', 2)]:
                with self.subTest(folder=folder):
                    output = base / folder
                    run = subprocess.run(command + [str(output)] + period_args, cwd=ROOT, capture_output=True, text=True)
                    self.assertEqual(run.returncode, expected_code, run.stderr)
                    self.assertTrue((output / 'relatorio.html').exists())
                    html = (output / 'relatorio.html').read_text(encoding='utf-8')
                    self.assertNotIn('<img src=x', html)
                    self.assertNotIn('Dados inteiramente fictícios', html)
                    self.assertNotIn('Setembro', html)
                    self.assertIn('Equipe &lt;img', html)
                    embedded = json.loads(re.search(r'id="demo-data">(.*?)</script>', html, re.S).group(1))
                    context = json.loads(re.search(r'id="report-context">(.*?)</script>', html, re.S).group(1))
                    self.assertEqual(context['mode'], 'run')
                    self.assertEqual(context['excelPaths'], {'execucao': 'fechamento.xlsx'})
                    self.assertEqual(embedded['execucao']['competencia'], '2026-03' if folder == 'pending' else '2026-02')
                    self.assertEqual(embedded['execucao']['resumo']['aprovadas'], 0 if folder == 'pending' else 2)
                    self.assertEqual(embedded['execucao'], json.loads((output / 'resultado.json').read_text(encoding='utf-8')))
                    if folder == 'out':
                        book = openpyxl.load_workbook(output / 'fechamento.xlsx', data_only=False)
                        self.assertEqual(book['Base aprovada']['G7'].value, 'FAT')
                        self.assertEqual(book['Base aprovada']['L7'].value, 'receita')
                        self.assertEqual(book['Base aprovada']['L8'].value, 'gasto')
                        book.close()
                        cached = openpyxl.load_workbook(output / 'fechamento.xlsx', data_only=True)
                        self.assertEqual(cached['Resumo']['B17'].value, 1250.50)
                        self.assertEqual(cached['Resumo']['B18'].value, 250.50)
                        self.assertEqual(cached['Resumo']['B19'].value, 1000)
                        cached.close()


if __name__ == '__main__':
    unittest.main()
