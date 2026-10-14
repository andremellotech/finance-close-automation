import csv
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from finance_close.engine import process, FIELDS
from finance_close.io import load_rows, load_controls
from finance_close.export import write_outputs
from test_engine import row, control, CATALOG

ROOT = Path(__file__).resolve().parents[1]


def write_csv(path, rows, headers=FIELDS):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


class PipelineTests(unittest.TestCase):
    def test_excel_error_cells_cannot_release_a_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "a.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Lancamentos"
            sheet.append(list(FIELDS))
            sheet.append(["#VALUE!", "2025-09-01", "ANA", "REV", "#N/A", 100])
            book.save(source)
            book.close()
            rows, sources = load_rows(directory)
            result = process(rows, control(1, "100", "a.xlsx"), "2025-09", CATALOG, sources)
            self.assertFalse(result["liberado"])
            self.assertEqual(result["linhas"][0]["status"], "PENDENTE")
            self.assertIn("ERRO_CELULA_EXCEL:registro_id", result["linhas"][0]["erros"])

    def test_uncached_trailing_formula_rows_are_traced(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "a.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Lancamentos"
            sheet.append(list(FIELDS))
            sheet.append(["A", "2025-09-01", "ANA", "REV", "primeiro", 100])
            sheet.append(['="B"', '="2025-09-01"', '="ANA"', '="REV"', '="segundo"', '=200'])
            book.save(source)
            book.close()
            rows, sources = load_rows(directory)
            self.assertEqual(len(rows), 2)
            result = process(rows, control(2, "300", "a.xlsx"), "2025-09", CATALOG, sources)
            self.assertEqual(result["linhas"][1]["linha"], 3)
            self.assertIn("FORMULA_NAO_ACEITA", result["linhas"][1]["erros"])

    def test_csv_multiline_fields_preserve_physical_source_line(self):
        with tempfile.TemporaryDirectory() as directory:
            write_csv(Path(directory) / "a.csv", [row(descricao="linha um\nlinha dois"), row("B", linha=4)])
            rows, _ = load_rows(directory)
            self.assertEqual([r["linha"] for r in rows], [2, 4])

    def test_nul_character_does_not_repair_invalid_money(self):
        with tempfile.TemporaryDirectory() as directory:
            write_csv(Path(directory) / "a.csv", [row(amount="100\x00invalid")])
            rows, sources = load_rows(directory)
            self.assertEqual(rows[0]["valor_brl"], "100\x00invalid")
            result = process(rows, control(), "2025-09", CATALOG, sources)
            self.assertFalse(result["liberado"])
            self.assertIn("VALOR_INVALIDO", result["linhas"][0]["erros"])

    def test_csv_reader_preserves_ids_and_sources(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "a.csv"
            write_csv(path, [row(id="0001")])
            before = path.read_bytes()
            rows, sources = load_rows(directory)
            self.assertEqual(rows[0]["registro_id"], "0001")
            self.assertEqual(rows[0]["linha"], 2)
            self.assertEqual(sources["a.csv"]["sha256"], hashlib.sha256(before).hexdigest())
            self.assertEqual(path.read_bytes(), before)

    def test_bad_headers_fail_with_filename(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalido.csv"
            path.write_text("id;valor\n1;10\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "invalido.csv"):
                load_rows(directory)

    def test_empty_header_only_file_does_not_release(self):
        with tempfile.TemporaryDirectory() as directory:
            write_csv(Path(directory) / "a.csv", [])
            rows, sources = load_rows(directory)
            result = process(rows, control(0, "0"), "2025-09", CATALOG, sources)
            self.assertFalse(result["liberado"])
            self.assertEqual(result["conciliacao"][0]["status"], "VAZIO")

    def test_duplicate_controls_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "controle.csv"
            path.write_text("arquivo;linhas_esperadas;total_esperado_brl\na.csv;1;100\na.csv;1;100\n", encoding="utf-8")
            with self.assertRaises(ValueError):
                load_controls(path)

    def test_exports_reconcile_and_formula_like_text_is_text(self):
        with tempfile.TemporaryDirectory() as directory:
            result = process([row(descricao="=1+1"), row("B", "-35.25", categoria="DIR", linha=3)], control(2, "64.75"), "2025-09", CATALOG)
            target = Path(directory) / "out"
            write_outputs(result, target)
            exported = json.loads((target / "resultado.json").read_text(encoding="utf-8"))
            self.assertEqual(exported["resumo"]["total_aprovado_centavos"], 6475)
            with (target / "base_aprovada.csv").open(encoding="utf-8-sig", newline="") as stream:
                records = list(csv.DictReader(stream, delimiter=";"))
            self.assertEqual(records[0]["descricao"], "'=1+1")
            book = openpyxl.load_workbook(target / "fechamento.xlsx", data_only=False)
            self.assertEqual(book.sheetnames, ["Resumo", "Base aprovada", "Pendencias", "Conciliacao", "Rastreabilidade"])
            self.assertEqual(book["Base aprovada"]["H7"].data_type, "s")
            self.assertEqual(book["Base aprovada"]["H7"].value, "=1+1")
            book.close()
            cached = openpyxl.load_workbook(target / "fechamento.xlsx", data_only=True)
            self.assertEqual(cached["Resumo"]["B19"].value, 64.75)
            self.assertEqual(cached["Resumo"]["B24"].value, 0)
            cached.close()

    def test_export_does_not_overwrite_existing_folder(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory)
            original = target / "meu-arquivo.txt"
            original.write_text("preservar", encoding="utf-8")
            result = process([row()], control(), "2025-09", CATALOG)
            with self.assertRaises(ValueError):
                write_outputs(result, target)
            self.assertEqual(original.read_text(), "preservar")

    def test_cli_exit_codes_and_input_safety(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            source = base / "input"
            source.mkdir()
            write_csv(source / "a.csv", [row()])
            (base / "controls.csv").write_text("arquivo;linhas_esperadas;total_esperado_brl\na.csv;1;100.00\n", encoding="utf-8")
            (base / "catalog.json").write_text(json.dumps(CATALOG), encoding="utf-8")
            command = [sys.executable, "-m", "finance_close", "run", "--input", str(source), "--controls", str(base / "controls.csv"), "--catalog", str(base / "catalog.json"), "--period", "2025-09", "--output"]
            good = subprocess.run(command + [str(base / "out")], cwd=ROOT, capture_output=True, text=True)
            self.assertEqual(good.returncode, 0, good.stderr)
            bad = subprocess.run(command + [str(base / "pending"), "--period", "2025-08"], cwd=ROOT, capture_output=True, text=True)
            self.assertEqual(bad.returncode, 2, bad.stderr)
            unsafe = subprocess.run(command + [str(source)], cwd=ROOT, capture_output=True, text=True)
            self.assertEqual(unsafe.returncode, 1)
            self.assertTrue((source / "a.csv").exists())


if __name__ == "__main__":
    unittest.main()
