"""Leitura de formatos públicos; nenhum arquivo de origem é escrito."""
import csv
import hashlib
import json
import io
from pathlib import Path

import pandas as pd
import openpyxl
from .engine import FIELDS
from .configuration import normalize_catalog, normalize_profile, normalize_record


def _headers(headers, expected, path):
    if len(set(headers)) != len(headers) or set(headers) != set(expected):
        raise ValueError(f"Cabeçalhos inválidos em {path.name}. Esperado: {', '.join(expected)}")


def _load_paths(files, profile=None):
    options = normalize_profile(profile)['entrada']
    expected = list(options['colunas'].values())

    def check_headers(headers, path):
        if len(set(headers)) != len(headers) or not set(expected) <= set(headers) or (not options['permitir_colunas_extras'] and set(headers) != set(expected)):
            raise ValueError(f"Cabeçalhos inválidos em {path.name}. Esperado pelo mapeamento: {', '.join(expected)}")
    rows, sources = [], {}
    for path in files:
        content = path.read_bytes()
        sources[path.name] = {"sha256": hashlib.sha256(content).hexdigest(), "bytes": len(content)}
        if path.suffix.lower() == ".csv":
            with io.StringIO(content.decode('utf-8-sig'), newline='') as stream:
                reader = csv.reader(stream, delimiter=options['separador'], strict=True)
                headers = next(reader, None)
                if headers is None:
                    raise ValueError("Arquivo sem cabeçalho: " + path.name)
                check_headers(headers, path)
                records, metadata = [], []
                previous_line = reader.line_num
                for values in reader:
                    line = previous_line + 1
                    previous_line = reader.line_num
                    if len(values) != len(headers):
                        raise ValueError(f"Quantidade de colunas inválida: {path.name}, linha {line}")
                    records.append(dict(zip(headers, values)))
                    metadata.append({"linha": line, "_formula": False, "_excel_errors": []})
            # Reuse the parsed values: a second parser could truncate NULs or infer IDs.
            frame = pd.DataFrame(records, columns=headers, dtype=object)
            sheet = "CSV"
        else:
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                if options['aba'] not in workbook.sheetnames:
                    raise ValueError(f"Aba {options['aba']} ausente em {path.name}")
                worksheet = workbook[options['aba']]
                headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
                check_headers(headers, path)
                records, metadata = [], []
                for cells in worksheet.iter_rows(min_row=2):
                    if all(cell.value is None for cell in cells):
                        continue
                    records.append(dict(zip(headers, [cell.value for cell in cells])))
                    selected = {field: cells[headers.index(column)] for field, column in options['colunas'].items()}
                    metadata.append({"linha": cells[0].row, "_formula": any(cell.data_type == "f" for cell in selected.values()), "_excel_errors": [field for field, cell in selected.items() if cell.data_type == "e"]})
            finally:
                workbook.close()
            # Preserve physical rows, raw formulas and error-cell types before conversion.
            frame = pd.DataFrame(records, columns=headers, dtype=object)
            sheet = options['aba']
        for values, origin in zip(frame.to_dict(orient="records"), metadata):
            rows.append({**normalize_record(values, options), "arquivo": path.name, "aba": sheet, **origin})
    return rows, sources


def _input_files(input_dir):
    directory = Path(input_dir)
    if not directory.is_dir():
        raise ValueError('A pasta de entrada não existe')
    files = sorted(p for p in directory.iterdir() if p.is_file() and p.suffix.lower() in {'.csv', '.xlsx'})
    if not files:
        raise ValueError('A entrada não contém arquivos CSV ou XLSX')
    return files


def load_rows(input_dir, profile=None):
    return _load_paths(_input_files(input_dir), profile)


def scan_rows(input_dir, profile=None):
    profile = normalize_profile(profile)
    rows, sources, issues = [], {}, []
    for path in _input_files(input_dir):
        try:
            current, source = _load_paths([path], profile)
            rows.extend(current)
            sources.update(source)
        except Exception as exc:
            issues.append({'codigo': 'LEITURA_INVALIDA', 'arquivo': path.name, 'problema': str(exc), 'acao': 'Confira se o arquivo está íntegro e legível, com aba, cabeçalhos, codificação e separador definidos no perfil.'})
    return rows, sources, issues


def load_controls(path):
    path = Path(path)
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, delimiter=";")
        _headers(reader.fieldnames or [], ("arquivo", "linhas_esperadas", "total_esperado_brl"), path)
        result = {}
        for row in reader:
            name = str(row["arquivo"] or "").strip()
            if not name or name in result or Path(name).name != name:
                raise ValueError("Controle com arquivo vazio, repetido ou caminho inválido")
            result[name] = row
        return result


def load_catalog(path):
    return normalize_catalog(json.loads(Path(path).read_text(encoding="utf-8")))
